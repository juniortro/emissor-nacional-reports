"""
Módulo de geração de relatórios Excel.
Exporta os dados coletados para arquivo .xlsx formatado.
"""
import logging
from datetime import datetime
from typing import List, Dict
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from config import Config

logger = logging.getLogger(__name__)


def generate_excel(notes: List[Dict], output_filename: str = None) -> str:
    """
    Gera arquivo Excel com os dados das notas fiscais.
    
    Args:
        notes: Lista de dicionários com dados das notas
        output_filename: Nome do arquivo (opcional, será gerado automaticamente se não fornecido)
    
    Returns:
        Caminho do arquivo gerado
    """
    if not notes:
        raise ValueError("Lista de notas vazia")
    
    logger.info(f"Gerando Excel com {len(notes)} notas...")
    
    # Gera nome do arquivo se não fornecido
    if not output_filename:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_filename = f"NFSe_Report_{timestamp}.xlsx"
    
    try:
        # Cria DataFrame
        df = pd.DataFrame(notes)
        
        # Reordena colunas para melhor visualização
        preferred_order = [
            'numero_nfse',
            'data_emissao',
            'situacao',
            'prestador_razao_social',
            'prestador_cnpj',
            'prestador_municipio',
            'tomador_razao_social',
            'tomador_cpf_cnpj',
            'tomador_municipio',
            'valor_servicos',
            'valor_deducoes',
            'base_calculo',
            'aliquota',
            'valor_iss',
            'iss_retido',
            'valor_pis',
            'valor_cofins',
            'valor_inss',
            'valor_ir',
            'valor_csll',
            'valor_liquido',
            'codigo_servico',
            'discriminacao_servicos',
            'codigo_verificacao'
        ]
        
        # Mantém apenas colunas que existem
        columns = [col for col in preferred_order if col in df.columns]
        # Adiciona colunas restantes que não estão na ordem preferida
        remaining_cols = [col for col in df.columns if col not in columns]
        columns.extend(remaining_cols)
        
        df = df[columns]
        
        # Renomeia colunas para nomes mais legíveis
        column_names = {
            'numero_nfse': 'Número NFSe',
            'data_emissao': 'Data Emissão',
            'situacao': 'Situação',
            'codigo_verificacao': 'Código Verificação',
            'prestador_razao_social': 'Prestador - Razão Social',
            'prestador_cnpj': 'Prestador - CNPJ',
            'prestador_municipio': 'Prestador - Município',
            'tomador_razao_social': 'Tomador - Razão Social',
            'tomador_cpf_cnpj': 'Tomador - CPF/CNPJ',
            'tomador_municipio': 'Tomador - Município',
            'valor_servicos': 'Valor Serviços',
            'valor_deducoes': 'Deduções',
            'base_calculo': 'Base Cálculo',
            'aliquota': 'Alíquota',
            'valor_iss': 'Valor ISS',
            'iss_retido': 'ISS Retido',
            'valor_pis': 'PIS',
            'valor_cofins': 'COFINS',
            'valor_inss': 'INSS',
            'valor_ir': 'IR',
            'valor_csll': 'CSLL',
            'valor_liquido': 'Valor Líquido',
            'codigo_servico': 'Código Serviço',
            'discriminacao_servicos': 'Discriminação dos Serviços'
        }
        
        df = df.rename(columns=column_names)
        
        # Salva Excel básico
        df.to_excel(output_filename, index=False, sheet_name='NFSe')
        
        # Aplica formatação
        format_excel(output_filename)
        
        logger.info(f"Relatório gerado com sucesso: {output_filename}")
        return output_filename
        
    except Exception as e:
        logger.error(f"Erro ao gerar Excel: {e}")
        raise


def format_excel(filename: str):
    """
    Aplica formatação ao arquivo Excel gerado.
    """
    try:
        wb = load_workbook(filename)
        ws = wb.active
        
        # Formatação do cabeçalho
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Aplica formatação ao cabeçalho
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
        
        # Ajusta largura das colunas
        column_widths = {
            'A': 15,  # Número NFSe
            'B': 15,  # Data Emissão
            'C': 12,  # Situação
            'D': 35,  # Prestador
            'E': 18,  # CNPJ Prestador
            'F': 20,  # Município Prestador
            'G': 35,  # Tomador
            'H': 18,  # CPF/CNPJ Tomador
            'I': 20,  # Município Tomador
            'J': 15,  # Valor Serviços
            'K': 12,  # Deduções
            'L': 15,  # Base Cálculo
            'M': 10,  # Alíquota
            'N': 12,  # ISS
            'O': 12,  # ISS Retido
            'P': 12,  # PIS
            'Q': 12,  # COFINS
            'R': 12,  # INSS
            'S': 12,  # IR
            'T': 12,  # CSLL
            'U': 15,  # Valor Líquido
            'V': 15,  # Código Serviço
            'W': 50,  # Discriminação
            'X': 25,  # Código Verificação
        }
        
        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width
        
        # Congela primeira linha
        ws.freeze_panes = 'A2'
        
        # Aplica filtro automático
        ws.auto_filter.ref = ws.dimensions
        
        # Alinhamento dos dados
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.column_letter in ['J', 'K', 'L', 'N', 'O', 'P', 'Q', 'R', 'S', 'T', 'U']:
                    # Valores monetários - alinhamento à direita
                    cell.alignment = Alignment(horizontal="right")
                else:
                    cell.alignment = Alignment(vertical="top", wrap_text=False)
        
        wb.save(filename)
        logger.info("Formatação aplicada com sucesso")
        
    except Exception as e:
        logger.warning(f"Erro ao formatar Excel (arquivo básico foi salvo): {e}")


def generate_csv(notes: List[Dict], output_filename: str = None) -> str:
    """
    Gera arquivo CSV com os dados das notas fiscais.
    
    Args:
        notes: Lista de dicionários com dados das notas
        output_filename: Nome do arquivo (opcional)
    
    Returns:
        Caminho do arquivo gerado
    """
    if not notes:
        raise ValueError("Lista de notas vazia")
    
    logger.info(f"Gerando CSV com {len(notes)} notas...")
    
    if not output_filename:
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        output_filename = f"NFSe_Report_{timestamp}.csv"
    
    try:
        df = pd.DataFrame(notes)
        df.to_csv(output_filename, index=False, encoding='utf-8-sig', sep=';')
        
        logger.info(f"CSV gerado com sucesso: {output_filename}")
        return output_filename
        
    except Exception as e:
        logger.error(f"Erro ao gerar CSV: {e}")
        raise
